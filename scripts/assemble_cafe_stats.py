"""T-08: 3개 브랜드를 한 번에 크롤링해서 {N월}_카페통계.xlsx 하나로 조립.

T-02(금호)/T-03(한국·넥센 공식 계정 필터링)/T-04(댓글 동일 배치)를 순서대로
모두 실행하고, PRD 3.1이 요구하는 시트 전부(금호타이어/한국타이어/넥센타이어/
금호/한국/넥센/댓글)를 하나의 워크북에 담는다. 스냅샷 시트("{N월} {일}일")는
2024년부터 누적되는 히스토리 표라 이전 달 파일에 이어붙이는 방식이 필요한데,
크롤러가 처음부터 만들 수 있는 게 아니라서(그리고 PRD 5장 완료 기준에도
없어서) 이번 스코프에서는 뺐다.
"""

from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import openpyxl  # noqa: E402

from naver_cafe_crawler.articles import write_brand_sheet  # noqa: E402
from naver_cafe_crawler.comments import write_comment_sheet  # noqa: E402
from naver_cafe_crawler.crawl import fetch_cafe_articles_and_comments_for_month  # noqa: E402
from naver_cafe_crawler.roster import load_raw  # noqa: E402
from naver_cafe_crawler.roster_sheet import write_roster_sheet  # noqa: E402

BRAND_SHEET_NAMES = {"kumho": "금호타이어", "hankook": "한국타이어", "nexen": "넥센타이어"}
ROSTER_SHEET_NAMES = {"kumho": "금호", "hankook": "한국", "nexen": "넥센"}
GUBUN = {"kumho": "금호", "hankook": "한국", "nexen": "넥센"}


def crawl_brand(brand_key: str, brand_data: dict, year: int, month: int, delay: float):
    official_accounts = brand_data.get("official_accounts") or []
    gubun = GUBUN[brand_key]
    all_cafes = [c for g in brand_data["groups"] for c in g["cafes"]]

    all_articles = []
    all_comments = []
    for i, cafe in enumerate(all_cafes, start=1):
        print(f"  [{brand_key} {i}/{len(all_cafes)}] {cafe['name']} ...", file=sys.stderr)
        try:
            articles, comments, restricted = fetch_cafe_articles_and_comments_for_month(
                cafe["url"],
                year,
                month,
                gubun=gubun,
                cafe_name=cafe["name"],
                group=cafe["car_model"],
                official_accounts=official_accounts,
            )
            all_articles.extend(articles)
            all_comments.extend(comments)
        except Exception as e:  # noqa: BLE001
            print(f"    [SKIP] {cafe['name']} 실패: {e}", file=sys.stderr)
        time.sleep(delay)
    return all_articles, all_comments


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--month", type=int, required=True)
    parser.add_argument("--config", default=str(Path(__file__).resolve().parents[1] / "config" / "cafes.yaml"))
    parser.add_argument("--out", required=True)
    parser.add_argument("--delay", type=float, default=0.12)
    args = parser.parse_args()

    data = load_raw(args.config)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    all_comments_by_brand = {}
    for brand_key in ["kumho", "hankook", "nexen"]:
        brand_data = data["brands"][brand_key]
        print(f"=== {brand_data['label']} 크롤링 시작 ===", file=sys.stderr)
        articles, comments = crawl_brand(brand_key, brand_data, args.year, args.month, args.delay)
        all_comments_by_brand[brand_key] = comments

        ws = wb.create_sheet(BRAND_SHEET_NAMES[brand_key])
        write_brand_sheet(ws, brand_data["label"], articles)

        roster_ws = wb.create_sheet(ROSTER_SHEET_NAMES[brand_key])
        write_roster_sheet(roster_ws, brand_data["label"], brand_data["groups"])

    all_comments = [c for comments in all_comments_by_brand.values() for c in comments]
    comment_ws = wb.create_sheet("댓글")
    write_comment_sheet(comment_ws, all_comments)

    wb.save(args.out)
    print(f"Wrote {args.out}", file=sys.stderr)


if __name__ == "__main__":
    main()
