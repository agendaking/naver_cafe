"""crawl_from_workbook.py가 만든 JSON을 운영현황 엑셀의 결과 시트에 기록한다.

- `금호타이어` 시트: 구분, 일자, 카페, 그룹, 게시글, 조회수, 좋아요, 댓글, 비고
- `한국타이어`/`넥센타이어` 시트: 구분, 일자, 그룹, 카페, 게시글, 조회수, 좋아요, 댓글, 비고
  (기존 양식의 컬럼 순서를 그대로 따름)
- 맨 뒤 `댓글` 시트: 금호타이어 게시판 게시글의 댓글 — 카페명, 아이디, 댓글내용, 작성일시

데이터는 5행(B열)부터, 게시판 리스트 순서 → 최신 게시글 순으로 쓴다 (8월 파일과 동일).
"""

from __future__ import annotations

import argparse
import json
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FIRST_ROW = 5
SHEETS = {"금호": "금호타이어", "한국": "한국타이어", "넥센": "넥센타이어"}
NUM_FMT = "#,##0_);[Red]\\(#,##0\\)"
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="맑은 고딕", size=11)


def style_cell(cell, *, number: bool = False) -> None:
    cell.font = copy(FONT)
    cell.border = copy(BORDER)
    if number:
        cell.number_format = NUM_FMT


def clear_rows(ws, start: int) -> None:
    for row in ws.iter_rows(min_row=start, min_col=2, max_col=10):
        for c in row:
            c.value = None


def write_brand(ws, brand_key: str, cafes: list[dict]) -> int:
    clear_rows(ws, FIRST_ROW)
    r = FIRST_ROW
    for cafe in cafes:
        for a in sorted(cafe["articles"], key=lambda x: x["datetime"], reverse=True):
            date = datetime.fromisoformat(a["datetime"]).strftime("%Y.%m.%d")
            if brand_key == "금호":
                cafe_col, group_col = 4, 5
            else:
                group_col, cafe_col = 4, 5
            values = {2: None, 3: date, cafe_col: cafe["name"], group_col: cafe["car_model"],
                      6: a["subject"], 7: a["views"], 8: a["likes"], 9: a["comments"], 10: None}
            for col, v in values.items():
                c = ws.cell(r, col, v)
                style_cell(c, number=col in (7, 8, 9))
            r += 1
    return r - FIRST_ROW


def write_comments(wb, cafes: list[dict]) -> int:
    if "댓글" in wb.sheetnames:
        del wb["댓글"]
    ws = wb.create_sheet("댓글")  # 맨 뒤에 추가됨
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 1.125
    for col, w in zip("BCDE", (24, 22, 80, 18)):
        ws.column_dimensions[col].width = w
    ws.cell(1, 2, "금호타이어 게시판 댓글 현황").font = Font(name="맑은 고딕", size=16, bold=True)
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    for i, label in enumerate(["카페명", "아이디", "댓글내용", "작성일시"], start=2):
        c = ws.cell(4, i, label)
        style_cell(c)
        c.font = Font(name="맑은 고딕", size=11, bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
    r = FIRST_ROW
    for cafe in cafes:
        for a in sorted(cafe["articles"], key=lambda x: x["datetime"], reverse=True):
            for cm in a["comment_items"] or []:
                vals = [cafe["name"], cm["writer"], cm["content"], datetime.fromisoformat(cm["datetime"])]
                for i, v in enumerate(vals, start=2):
                    c = ws.cell(r, i, v)
                    style_cell(c)
                ws.cell(r, 5).number_format = "yyyy-mm-dd hh:mm"
                ws.cell(r, 4).alignment = Alignment(wrap_text=False)
                r += 1
    ws.freeze_panes = "B5"
    return r - FIRST_ROW


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, help="원본 운영현황 엑셀")
    parser.add_argument("--data", required=True, help="crawl_from_workbook.py 출력 JSON")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    data = json.loads(open(args.data, encoding="utf-8").read())
    wb = openpyxl.load_workbook(args.workbook)
    for key, sheet in SHEETS.items():
        cafes = data["brands"][key]["cafes"]
        n = write_brand(wb[sheet], key, cafes)
        arts = [a for c in cafes for a in c["articles"]]
        print(f"{sheet}: {n}건 / 조회수 {sum(a['views'] for a in arts):,} "
              f"좋아요 {sum(a['likes'] for a in arts):,} 댓글 {sum(a['comments'] for a in arts):,}")
    n = write_comments(wb, data["brands"]["금호"]["cafes"])
    print(f"댓글: {n}건")
    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
