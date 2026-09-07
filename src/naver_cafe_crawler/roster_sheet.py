"""카페 로스터 시트(PRD 3.1 `금호`/`한국`/`넥센` 시트) 작성.

config/cafes.yaml의 구조화된 로스터를 그대로 xlsx 시트로 옮긴다. 원본
샘플 파일들은 브랜드마다 컬럼 배치가 조금씩 다르고 forward-fill에
의존했는데(T-01에서 확인한 파싱 함정), 여기서는 세 브랜드 모두 같은
고정 컬럼으로, 매 행에 값을 채워서 쓴다.
"""

from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

ROSTER_HEADER = ["No", "게시판그룹", "차종", "카페명", "전체회원수", "URL"]


def write_roster_sheet(ws: Worksheet, brand_label: str, groups: list[dict]) -> None:
    ws.cell(1, 2, f"{brand_label} 자동차 동호회 리스트")

    for col_idx, label in enumerate(ROSTER_HEADER, start=2):
        ws.cell(3, col_idx, label)

    row = 4
    no = 1
    for group in groups:
        for cafe in group["cafes"]:
            ws.cell(row, 2, no)
            ws.cell(row, 3, group["board_group"])
            ws.cell(row, 4, cafe["car_model"])
            ws.cell(row, 5, cafe["name"])
            ws.cell(row, 6, cafe.get("member_count"))
            ws.cell(row, 7, cafe["url"])
            row += 1
            no += 1
