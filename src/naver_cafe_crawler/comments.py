"""댓글 원문 시트(PRD 3.1 `댓글` 시트) 작성.

이 시트는 반드시 게시글 데이터(articles.py)와 같은 크롤링 배치에서 채워야
한다 — 별도로 돌리면 수집 시점 차이 때문에 총량이 안 맞는 문제가 실제로
있었다 (PRD 3.1). crawl.py의 fetch_cafe_articles_and_comments_for_month()가
이걸 구조적으로 보장한다 (게시글 목록을 한 번만 가져와서 그 결과로 댓글도
같이 가져옴).
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.worksheet.worksheet import Worksheet

# PRD 3.1: 댓글 시트 컬럼 순서.
COMMENT_HEADER = ["구분", "카페", "게시물", "발행일", "작성자", "댓글내용"]


@dataclass(frozen=True)
class CommentRecord:
    gubun: str  # 구분: 금호/한국/넥센
    cafe: str
    article_title: str
    write_date: str  # "2026.08.28"
    writer: str
    content: str


def write_comment_sheet(ws: Worksheet, comments: list[CommentRecord]) -> None:
    for col_idx, label in enumerate(COMMENT_HEADER, start=1):
        ws.cell(1, col_idx, label)

    for row_offset, c in enumerate(comments):
        r = 2 + row_offset
        ws.cell(r, 1, c.gubun)
        ws.cell(r, 2, c.cafe)
        ws.cell(r, 3, c.article_title)
        ws.cell(r, 4, c.write_date)
        ws.cell(r, 5, c.writer)
        ws.cell(r, 6, c.content)
