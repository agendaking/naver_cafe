"""브랜드별 게시글 현황 시트(PRD 3.1 `금호타이어`/`한국타이어`/`넥센타이어` 시트)를
만드는 공통 로직.

세 브랜드 시트의 컬럼 순서가 어긋나면 안 된다는 게 PRD가 가장 강하게 못박은
요건이라, 컬럼 순서를 브랜드마다 따로 작성하지 않고 이 모듈 하나로 고정한다
(T-02가 금호를 쓰고, T-03이 한국/넥센에 그대로 재사용).
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass

from openpyxl.worksheet.worksheet import Worksheet

# PRD 3.1: 세 브랜드 시트 모두 이 순서로 고정.
HEADER = ["구분", "일자", "카페", "그룹", "게시글", "조회수", "좋아요", "댓글", "비고"]


@dataclass(frozen=True)
class ArticleRecord:
    date: str  # "2026.08.28" 형식
    cafe: str
    group: str  # 차종
    title: str
    views: int
    likes: int
    comments: int
    gubun: str | None = None  # 구분 (필요 시에만 채움, 원본도 대부분 비어있음)
    note: str | None = None  # 비고


def normalize_title(title: str) -> str:
    """크로스포스팅된 동일 콘텐츠를 같은 "콘텐츠"로 식별하기 위한 제목 정규화.

    - 앞뒤 공백 제거, 연속 공백/탭을 단일 공백으로
    - 유니코드 NFKC 정규화 (전각/반각, 장식용 수학 알파벳 등을 표준 문자로 변환.
      예: 'ℂ' -> 'C', '𝕓𝕪' -> 'by')
    - 카페마다 붙이는 게시판 태그(예: "[공지]", "[이벤트]")는 내용 자체가 아니므로 제거

    완전히 동일한 정규화 결과가 나오는 게시글들만 "같은 콘텐츠"로 취급한다.
    카페별로 제목을 아예 다르게 쓴 경우까지 잡아내지는 않는다 (범위 밖).
    """
    normalized = unicodedata.normalize("NFKC", title)
    normalized = re.sub(r"^\s*\[[^\]]*\]\s*", "", normalized)
    normalized = re.sub(r"\s+", " ", normalized).strip()
    return normalized


def group_cross_posted(articles: list[ArticleRecord]) -> dict[str, list[ArticleRecord]]:
    """정규화된 제목 기준으로 여러 카페에 걸쳐 올라온 동일 콘텐츠를 묶는다.

    출력 시트 자체는 병합하지 않고 게시글별로 한 행씩 그대로 쓴다 (PRD 3.1) —
    이 함수는 콘텐츠 식별/집계용 부가 정보를 만들 뿐이다.
    """
    groups: dict[str, list[ArticleRecord]] = {}
    for article in articles:
        key = normalize_title(article.title)
        groups.setdefault(key, []).append(article)
    return groups


def write_brand_sheet(ws: Worksheet, brand_label: str, articles: list[ArticleRecord]) -> None:
    """PRD 3.1 스키마대로 게시글 현황 시트를 채운다.

    row1: 타이틀, row3: 합계, row4: 헤더, row5부터: 데이터.
    합계 행의 조회수/좋아요/댓글은 데이터 행 합과 항상 정확히 일치한다
    (SUM 수식이 아니라 파이썬에서 직접 더한 값을 쓴다 — 엑셀에서 셀을 열어보지
    않아도, 값 자체로 이미 검증 가능하게).
    """
    ws.cell(1, 2, f"{brand_label} 게시글 현황")

    total_views = sum(a.views for a in articles)
    total_likes = sum(a.likes for a in articles)
    total_comments = sum(a.comments for a in articles)

    ws.cell(3, 2, "합계")
    ws.cell(3, 3, "-")
    ws.cell(3, 4, "-")
    ws.cell(3, 5, None)
    ws.cell(3, 6, "-")
    ws.cell(3, 7, total_views)
    ws.cell(3, 8, total_likes)
    ws.cell(3, 9, total_comments)

    for col_idx, label in enumerate(HEADER, start=2):
        ws.cell(4, col_idx, label)

    for row_offset, article in enumerate(articles):
        r = 5 + row_offset
        ws.cell(r, 2, article.gubun)
        ws.cell(r, 3, article.date)
        ws.cell(r, 4, article.cafe)
        ws.cell(r, 5, article.group)
        ws.cell(r, 6, article.title)
        ws.cell(r, 7, article.views)
        ws.cell(r, 8, article.likes)
        ws.cell(r, 9, article.comments)
        ws.cell(r, 10, article.note)
