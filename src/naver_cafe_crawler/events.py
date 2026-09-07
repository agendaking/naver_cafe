"""이벤트 참여 현황 (PRD 3.2).

형님이 매달 이벤트 게시글 제목(또는 그 중 한 카페 URL)을 직접 지정하면,
그 게시글이 크로스포스팅된 모든 금호 카페에서 조회수/좋아요/댓글과 댓글
작성자 목록을 모은다.

"응모"/"추첨대상" 숫자는 정확한 필터링 규칙(중복 참여 제외, 자격 제외 등)이
PRD/TASKS 어디에도 정의돼 있지 않아 추측하지 않는다. 대신:
- 조회수/좋아요/댓글: 실측 검증 가능한 원본 수치 그대로
- 고유 참여자수: 댓글 작성자 닉네임을 카페 간 중복 제거한 수 (참고용 근사치.
  실제 "응모" 집계 규칙과 다를 수 있음 — 형님 확인 필요)
"""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.worksheet.worksheet import Worksheet

from .articles import normalize_title
from .comments import CommentRecord
from .crawl import (
    FetchCommentPageFn,
    FetchPageFn,
    fetch_comments_for_article,
    fetch_raw_articles_for_month,
)
from .naver_api import (
    CommentsRestrictedError,
    fetch_article_page,
    fetch_comment_page,
    parse_cafe_menu_ids,
)

EVENT_HEADER = [
    "No", "게시판그룹", "차종", "카페명", "URL", "조회수", "좋아요", "댓글", "발견여부",
]


@dataclass(frozen=True)
class EventCafeResult:
    no: int
    board_group: str
    car_model: str
    cafe_name: str
    url: str
    views: int
    likes: int
    comments: int
    found: bool
    article_id: int | None = None
    cafe_id: str | None = None


@dataclass(frozen=True)
class EventSummary:
    total_views: int
    total_likes: int
    total_comments: int
    unique_participant_count: int
    cafes_found_in: int
    cafes_total: int
    restricted_comment_count: int = 0


def find_event_across_cafes(
    cafes: list[dict],
    event_title: str,
    year: int,
    month: int,
    *,
    fetch_page: FetchPageFn = fetch_article_page,
) -> list[EventCafeResult]:
    """카페 로스터 전체를 대상 월로 훑어서, 제목이 일치하는 게시글을 찾는다."""
    target_key = normalize_title(event_title)
    results: list[EventCafeResult] = []

    for i, cafe in enumerate(cafes, start=1):
        cafe_id, menu_id = parse_cafe_menu_ids(cafe["url"])
        raw = fetch_raw_articles_for_month(cafe_id, menu_id, year, month, fetch_page=fetch_page)
        match = next((a for a in raw if normalize_title(a.subject) == target_key), None)
        if match is None:
            results.append(
                EventCafeResult(
                    no=i,
                    board_group=cafe.get("board_group", ""),
                    car_model=cafe["car_model"],
                    cafe_name=cafe["name"],
                    url=cafe["url"],
                    views=0,
                    likes=0,
                    comments=0,
                    found=False,
                )
            )
            continue
        results.append(
            EventCafeResult(
                no=i,
                board_group=cafe.get("board_group", ""),
                car_model=cafe["car_model"],
                cafe_name=cafe["name"],
                url=cafe["url"],
                views=match.read_count,
                likes=match.like_count,
                comments=match.comment_count,
                found=True,
                article_id=match.article_id,
                cafe_id=cafe_id,
            )
        )
    return results


def collect_event_comments(
    results: list[EventCafeResult],
    *,
    gubun: str,
    fetch_comment_page: FetchCommentPageFn = fetch_comment_page,
) -> tuple[list[CommentRecord], int]:
    """찾은 게시글들의 댓글(=참여자 목록)을 전부 모은다.

    반환값의 두 번째 원소는 회원전용이라 못 가져온 댓글 수 합계 (T-04와 동일한
    예외 처리).
    """
    all_comments: list[CommentRecord] = []
    restricted = 0
    for r in results:
        if not r.found or r.comments == 0:
            continue
        try:
            all_comments.extend(
                fetch_comments_for_article(
                    r.cafe_id,
                    r.article_id,
                    gubun=gubun,
                    cafe_name=r.cafe_name,
                    article_title=r.cafe_name,
                    fetch_comment_page=fetch_comment_page,
                )
            )
        except CommentsRestrictedError:
            restricted += r.comments
    return all_comments, restricted


def summarize(results: list[EventCafeResult], comments: list[CommentRecord], restricted: int) -> EventSummary:
    found = [r for r in results if r.found]
    unique_participants = {c.writer for c in comments}
    return EventSummary(
        total_views=sum(r.views for r in found),
        total_likes=sum(r.likes for r in found),
        total_comments=sum(r.comments for r in found),
        unique_participant_count=len(unique_participants),
        cafes_found_in=len(found),
        cafes_total=len(results),
        restricted_comment_count=restricted,
    )


def write_event_sheet(
    ws: Worksheet, event_title: str, results: list[EventCafeResult], summary: EventSummary
) -> None:
    ws.cell(1, 1, f"이벤트: {event_title}")
    ws.cell(2, 1, "총조회수")
    ws.cell(2, 2, "총좋아요")
    ws.cell(2, 3, "총댓글")
    ws.cell(2, 4, "고유 참여자수(근사치)")
    ws.cell(2, 5, "발견된 카페 수")
    ws.cell(2, 6, "회원전용 제외 댓글수")
    ws.cell(3, 1, summary.total_views)
    ws.cell(3, 2, summary.total_likes)
    ws.cell(3, 3, summary.total_comments)
    ws.cell(3, 4, summary.unique_participant_count)
    ws.cell(3, 5, f"{summary.cafes_found_in}/{summary.cafes_total}")
    ws.cell(3, 6, summary.restricted_comment_count)

    header_row = 5
    for col_idx, label in enumerate(EVENT_HEADER, start=1):
        ws.cell(header_row, col_idx, label)

    for row_offset, r in enumerate(results):
        row = header_row + 1 + row_offset
        ws.cell(row, 1, r.no)
        ws.cell(row, 2, r.board_group)
        ws.cell(row, 3, r.car_model)
        ws.cell(row, 4, r.cafe_name)
        ws.cell(row, 5, r.url)
        ws.cell(row, 6, r.views)
        ws.cell(row, 7, r.likes)
        ws.cell(row, 8, r.comments)
        ws.cell(row, 9, "O" if r.found else "X")
