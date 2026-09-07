"""카카오톡 채널 통계 (PRD 3.3).

이 데이터는 크롤링이 아니라 형님이 매달 카카오톡 채널 관리자센터
(center-pf.kakao.com)에서 직접 다운로드해오는 xls 파일을 입력으로 받는다
(형님 확인, 2026-09-07). 이 모듈이 하는 일은:
1. 원본 xls를 읽어서 구조화된 레코드로 파싱
2. 대상 월의 날짜가 하루도 빠짐없이, 중복 없이 들어있는지 검증
3. PRD 3.3 스키마 그대로 xlsx로 저장

수치 자체가 관리자센터 화면과 맞는지는 형님이 직접 대조해야 하는 부분이라
(승인 기준), 이 모듈은 "형식이 깨지지 않았는지"까지만 자동 검증한다.
"""

from __future__ import annotations

from calendar import monthrange
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook

FRIEND_SHEET_NAME = "친구수 통계"
VIEW_SHEET_NAME = "조회수와 방문자수"

FRIEND_HEADER = ["날짜", "친구수", "채널 추가수 합계", "채팅 요청 친구수"]
VIEW_HEADER = ["날짜", "방문자수", "조회수"]
MESSAGE_HEADER = ["발송횟수", "발송건수", "발송비용", "카카오캐시잔액"]


@dataclass(frozen=True)
class FriendStatRow:
    date: str  # "2026-08-31"
    friend_count: int
    channel_add_count: int
    chat_request_count: int


@dataclass(frozen=True)
class ViewStatRow:
    date: str
    visitor_count: int
    view_count: int


@dataclass(frozen=True)
class MessageSummary:
    send_count: int
    send_volume: int
    send_cost: int
    kakao_cash_balance: int


@dataclass(frozen=True)
class KakaoData:
    friends: list[FriendStatRow]
    message_summary: MessageSummary | None
    views: list[ViewStatRow]


def _read_xls_sheet_rows(path: str | Path, sheet_name: str) -> list[list]:
    import xlrd

    wb = xlrd.open_workbook(path)
    ws = wb.sheet_by_name(sheet_name)
    return [ws.row_values(r) for r in range(ws.nrows)]


def parse_friend_rows(rows: list[list]) -> tuple[list[FriendStatRow], MessageSummary | None]:
    """친구수 통계 시트의 원본 행들(헤더 포함)을 파싱한다.

    '메시지 발송' 요약 블록은 G~J열(인덱스 6~9) 어딘가에 3행짜리 미니 표로
    끼어 있다 (라벨행 -> 헤더행 -> 값행). 어느 날짜 행과 나란히 있는지는
    의미가 없으므로, 위치를 가정하지 않고 라벨을 찾아서 파싱한다.
    """
    friends: list[FriendStatRow] = []
    message_summary: MessageSummary | None = None

    for i, row in enumerate(rows):
        if i == 0:
            continue  # 헤더 행
        date = row[0]
        if date:
            friends.append(
                FriendStatRow(
                    date=str(date),
                    friend_count=int(row[1]) if row[1] != "" else 0,
                    channel_add_count=int(row[2]) if row[2] != "" else 0,
                    chat_request_count=int(row[3]) if row[3] != "" else 0,
                )
            )
        if len(row) > 6 and isinstance(row[6], str) and "메시지 발송" in row[6]:
            header_row = rows[i + 1]
            value_row = rows[i + 2]
            if header_row[6:10] != MESSAGE_HEADER:
                raise ValueError(f"메시지 발송 블록 헤더가 예상과 다름: {header_row[6:10]}")
            send_count, send_volume, send_cost, cash_balance = value_row[6:10]
            message_summary = MessageSummary(
                send_count=int(send_count),
                send_volume=int(send_volume),
                send_cost=int(send_cost),
                kakao_cash_balance=int(cash_balance),
            )

    return friends, message_summary


def parse_view_rows(rows: list[list]) -> list[ViewStatRow]:
    views: list[ViewStatRow] = []
    for i, row in enumerate(rows):
        if i == 0:
            continue
        if not row[0]:
            continue
        views.append(
            ViewStatRow(date=str(row[0]), visitor_count=int(row[1]), view_count=int(row[2]))
        )
    return views


def load_kakao_xls(path: str | Path) -> KakaoData:
    friend_rows = _read_xls_sheet_rows(path, FRIEND_SHEET_NAME)
    view_rows = _read_xls_sheet_rows(path, VIEW_SHEET_NAME)
    friends, message_summary = parse_friend_rows(friend_rows)
    views = parse_view_rows(view_rows)
    return KakaoData(friends=friends, message_summary=message_summary, views=views)


def validate_month_completeness(dates: list[str], year: int, month: int) -> list[str]:
    """대상 월의 날짜가 하루도 안 빠지고, 중복도 없이 들어있는지 확인.

    반환값이 빈 리스트면 통과. 아니면 문제를 설명하는 메시지 목록.
    """
    days_in_month = monthrange(year, month)[1]
    expected = {f"{year:04d}-{month:02d}-{d:02d}" for d in range(1, days_in_month + 1)}
    actual = set(dates)

    errors = []
    missing = sorted(expected - actual)
    if missing:
        errors.append(f"누락된 날짜: {missing}")
    extra = sorted(actual - expected)
    if extra:
        errors.append(f"대상 월({year}-{month:02d}) 밖의 날짜: {extra}")
    dup_counts = Counter(dates)
    dups = sorted(d for d, c in dup_counts.items() if c > 1)
    if dups:
        errors.append(f"중복된 날짜: {dups}")
    return errors


def write_kakao_workbook(out_path: str | Path, data: KakaoData) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = FRIEND_SHEET_NAME
    ws1.append(FRIEND_HEADER)
    for f in data.friends:
        ws1.append([f.date, f.friend_count, f.channel_add_count, f.chat_request_count])

    if data.message_summary is not None:
        m = data.message_summary
        # 원본 관례대로 두 번째 데이터 행(row 3) 옆에 라벨을 둔다. 어느 행과
        # 나란히 있든 의미는 없다 (parse_friend_rows도 위치를 안 가정함).
        ws1.cell(3, 7, "메시지 발송")
        ws1.cell(4, 7, MESSAGE_HEADER[0])
        ws1.cell(4, 8, MESSAGE_HEADER[1])
        ws1.cell(4, 9, MESSAGE_HEADER[2])
        ws1.cell(4, 10, MESSAGE_HEADER[3])
        ws1.cell(5, 7, m.send_count)
        ws1.cell(5, 8, m.send_volume)
        ws1.cell(5, 9, m.send_cost)
        ws1.cell(5, 10, m.kakao_cash_balance)

    ws2 = wb.create_sheet(VIEW_SHEET_NAME)
    ws2.append(VIEW_HEADER)
    for v in data.views:
        ws2.append([v.date, v.visitor_count, v.view_count])

    wb.save(out_path)
