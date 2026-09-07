"""T-05: 카카오톡 채널 통계 파싱/검증/변환 테스트.

승인 기준(카카오톡 채널 관리자센터 화면과 수치 대조)은 형님이 직접 확인해야
하는 부분이라 자동화할 수 없다. 이 테스트는 그 전 단계, 즉 원본 xls 구조를
깨지 않고 파싱해서 PRD 3.3 스키마로 정확히 옮기는지와, 날짜 누락/중복을
잡아내는지를 검증한다.
"""

from __future__ import annotations

import openpyxl
import pytest

from naver_cafe_crawler.kakao import (
    FRIEND_HEADER,
    MESSAGE_HEADER,
    VIEW_HEADER,
    KakaoData,
    parse_friend_rows,
    parse_view_rows,
    validate_month_completeness,
    write_kakao_workbook,
)

FRIEND_HEADER_ROW = ["날짜", "친구수", "채널 추가수 합계", "채팅 요청 친구수", "", "", "", "", "", ""]


def test_parse_friend_rows_basic():
    rows = [
        FRIEND_HEADER_ROW,
        ["2026-08-31", 10064, 1, 1, "", "", "", "", "", ""],
        ["2026-08-30", 10064, 0, 1, "", "", "", "", "", ""],
    ]
    friends, summary = parse_friend_rows(rows)
    assert len(friends) == 2
    assert friends[0].date == "2026-08-31"
    assert friends[0].friend_count == 10064
    assert summary is None


def test_parse_friend_rows_finds_message_summary_wherever_it_is():
    # 메시지 발송 블록이 어느 날짜 행 옆에 있든 상관없이 파싱되어야 함.
    rows = [
        FRIEND_HEADER_ROW,
        ["2026-08-31", 10064, 1, 1, "", "", "", "", "", ""],
        ["2026-08-30", 10064, 0, 1, "", "", "8월 메시지 발송", "", "", ""],
        ["2026-08-29", 10066, 2, 3, "", "", "발송횟수", "발송건수", "발송비용", "카카오캐시잔액"],
        ["2026-08-28", 10065, 0, 2, "", "", 1, 10175, 167887, 2944202],
    ]
    friends, summary = parse_friend_rows(rows)
    assert len(friends) == 4
    assert summary is not None
    assert summary.send_count == 1
    assert summary.send_volume == 10175
    assert summary.send_cost == 167887
    assert summary.kakao_cash_balance == 2944202


def test_parse_friend_rows_bad_message_header_raises():
    rows = [
        FRIEND_HEADER_ROW,
        ["2026-08-31", 10064, 1, 1, "", "", "메시지 발송", "", "", ""],
        ["2026-08-30", 10064, 0, 1, "", "", "엉뚱한헤더", "", "", ""],
        ["2026-08-29", 10066, 2, 3, "", "", 1, 2, 3, 4],
    ]
    with pytest.raises(ValueError, match="메시지 발송 블록 헤더"):
        parse_friend_rows(rows)


def test_parse_view_rows():
    rows = [
        ["날짜", "방문자수", "조회수"],
        ["2026-08-31", 3, 4],
        ["2026-08-30", 3, 7],
    ]
    views = parse_view_rows(rows)
    assert len(views) == 2
    assert views[0].visitor_count == 3
    assert views[0].view_count == 4


def test_validate_month_completeness_detects_missing_day():
    dates = [f"2026-08-{d:02d}" for d in range(1, 31)]  # 31일 빠짐
    errors = validate_month_completeness(dates, 2026, 8)
    assert any("누락" in e for e in errors)
    assert "2026-08-31" in errors[0]


def test_validate_month_completeness_detects_duplicate():
    dates = [f"2026-08-{d:02d}" for d in range(1, 32)] + ["2026-08-15"]
    errors = validate_month_completeness(dates, 2026, 8)
    assert any("중복" in e for e in errors)


def test_validate_month_completeness_detects_out_of_range():
    dates = [f"2026-08-{d:02d}" for d in range(1, 32)]
    dates[0] = "2026-09-01"
    errors = validate_month_completeness(dates, 2026, 8)
    assert any("밖" in e for e in errors)


def test_validate_month_completeness_passes_for_full_month():
    dates = [f"2026-08-{d:02d}" for d in range(1, 32)]
    assert validate_month_completeness(dates, 2026, 8) == []


def test_write_kakao_workbook_roundtrip(tmp_path):
    rows = [
        FRIEND_HEADER_ROW,
        ["2026-08-31", 10064, 1, 1, "", "", "", "", "", ""],
        ["2026-08-30", 10064, 0, 1, "", "", "8월 메시지 발송", "", "", ""],
        ["2026-08-29", 10066, 2, 3, "", "", "발송횟수", "발송건수", "발송비용", "카카오캐시잔액"],
        ["2026-08-28", 10065, 0, 2, "", "", 1, 10175, 167887, 2944202],
    ]
    friends, summary = parse_friend_rows(rows)
    view_rows = [["날짜", "방문자수", "조회수"], ["2026-08-31", 3, 4]]
    views = parse_view_rows(view_rows)

    data = KakaoData(friends=friends, message_summary=summary, views=views)
    out = tmp_path / "kakao.xlsx"
    write_kakao_workbook(out, data)

    wb = openpyxl.load_workbook(out, data_only=True)
    assert wb.sheetnames == ["친구수 통계", "조회수와 방문자수"]

    ws1 = wb["친구수 통계"]
    assert [ws1.cell(1, c).value for c in range(1, 5)] == FRIEND_HEADER
    assert [ws1.cell(2, c).value for c in range(1, 5)] == ["2026-08-31", 10064, 1, 1]
    assert [ws1.cell(4, c).value for c in range(7, 11)] == MESSAGE_HEADER
    assert [ws1.cell(5, c).value for c in range(7, 11)] == [1, 10175, 167887, 2944202]

    ws2 = wb["조회수와 방문자수"]
    assert [ws2.cell(1, c).value for c in range(1, 4)] == VIEW_HEADER
    assert [ws2.cell(2, c).value for c in range(1, 4)] == ["2026-08-31", 3, 4]
