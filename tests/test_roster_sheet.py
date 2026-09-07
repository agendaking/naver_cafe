"""T-08 승인 기준 검증(카페명 1:1 매칭)에 쓰일 로스터 시트 작성 테스트."""

from __future__ import annotations

import openpyxl

from naver_cafe_crawler.roster_sheet import ROSTER_HEADER, write_roster_sheet

GROUPS = [
    {
        "board_group": "CRUGEN GT Pro",
        "cafes": [
            {"car_model": "카니발KA4", "name": "카니발 포에버", "url": "https://cafe.naver.com/f-e/cafes/1/menus/1", "member_count": 699138},
        ],
    },
    {
        "board_group": "금호타이어",
        "cafes": [
            {"car_model": "쏘나타DN8", "name": "쏘나타 오너스클럽", "url": "https://cafe.naver.com/f-e/cafes/2/menus/2", "member_count": 537292},
        ],
    },
]


def test_write_roster_sheet_schema(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "금호"
    write_roster_sheet(ws, "금호타이어", GROUPS)
    out = tmp_path / "roster.xlsx"
    wb.save(out)

    loaded = openpyxl.load_workbook(out, data_only=True)["금호"]
    header = [loaded.cell(3, c).value for c in range(2, 2 + len(ROSTER_HEADER))]
    assert header == ROSTER_HEADER
    assert loaded.cell(4, 2).value == 1
    assert loaded.cell(4, 3).value == "CRUGEN GT Pro"
    assert loaded.cell(4, 5).value == "카니발 포에버"
    assert loaded.cell(5, 2).value == 2
    assert loaded.cell(5, 5).value == "쏘나타 오너스클럽"


def test_write_roster_sheet_numbers_continuously_across_groups():
    wb = openpyxl.Workbook()
    ws = wb.active
    write_roster_sheet(ws, "금호타이어", GROUPS)
    nos = [ws.cell(r, 2).value for r in range(4, 6)]
    assert nos == [1, 2]
