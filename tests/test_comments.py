"""T-04 승인 기준 검증:
1. 댓글 시트 스키마가 PRD 3.1대로 나오는지
2. 댓글이 게시글과 같은 배치(같은 fetch 호출)에서 나오는지, 그리고
   `댓글` 시트 행 수(구분 필터 후)가 게시글 시트의 댓글 합계와 정확히
   일치하는지
"""

from __future__ import annotations

import openpyxl
import pytest

from naver_cafe_crawler.comments import (
    COMMENT_HEADER,
    FIRST_DATA_ROW,
    HEADER_ROW,
    CommentRecord,
    write_comment_sheet,
)
from naver_cafe_crawler.crawl import (
    fetch_cafe_articles_and_comments_for_month,
    fetch_comments_for_article,
)
from naver_cafe_crawler.naver_api import CommentsRestrictedError

SAMPLE_COMMENTS = [
    CommentRecord(gubun="금호", cafe="카니발 포에버", article_title="휴가철 이벤트",
                  write_date="2026.08.03", writer="헌혈500", content="오! 좋네요"),
    CommentRecord(gubun="금호", cafe="카니발 포에버", article_title="휴가철 이벤트",
                  write_date="2026.08.04", writer="타이어마니아", content="참여합니다"),
]


@pytest.fixture
def sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "댓글"
    write_comment_sheet(ws, SAMPLE_COMMENTS)
    out = tmp_path / "comments.xlsx"
    wb.save(out)
    return openpyxl.load_workbook(out, data_only=True)["댓글"]


def test_comment_header_matches_prd_order(sheet):
    header = [sheet.cell(HEADER_ROW, c).value for c in range(1, 1 + len(COMMENT_HEADER))]
    assert header == COMMENT_HEADER


def test_comment_rows_written(sheet):
    assert sheet.cell(FIRST_DATA_ROW, 1).value == "금호"
    assert sheet.cell(FIRST_DATA_ROW, 2).value == "카니발 포에버"
    assert sheet.cell(FIRST_DATA_ROW, 6).value == "오! 좋네요"
    assert sheet.cell(FIRST_DATA_ROW + 1, 6).value == "참여합니다"


# --- 동일 배치 수집 검증 ---

def _article_page(cafe_id, menu_id, page, **kwargs):
    assert page == 1
    return {
        "result": {
            "articleList": [
                {
                    "type": "ARTICLE",
                    "item": {
                        "articleId": 111,
                        "subject": "댓글 2개짜리 글",
                        "writeDateTimestamp": 1786501304010,  # 2026-08-12
                        "readCount": 10,
                        "likeCount": 1,
                        "commentCount": 2,
                        "writerInfo": {"nickName": "타이어프로", "memberLevelName": "협력업체"},
                    },
                },
                {
                    "type": "ARTICLE",
                    "item": {
                        "articleId": 112,
                        "subject": "댓글 0개짜리 글",
                        "writeDateTimestamp": 1786673637613,  # 2026-08-14
                        "readCount": 5,
                        "likeCount": 0,
                        "commentCount": 0,
                        "writerInfo": {"nickName": "타이어프로", "memberLevelName": "협력업체"},
                    },
                },
            ],
            "pageInfo": {"lastNavigationPageNumber": 1, "visibleNextButton": False},
        }
    }


def _comment_page_for_111(cafe_id, article_id, page, **kwargs):
    assert article_id == 111  # 댓글 0개인 112번 글은 애초에 호출되면 안 됨
    return {
        "result": {
            "comments": {
                "items": [
                    {"writer": {"nick": "A"}, "content": "첫댓글", "updateDate": 1786501400000, "isDeleted": False},
                    {"writer": {"nick": "B"}, "content": "둘째댓글", "updateDate": 1786501500000, "isDeleted": False},
                ]
            },
            "hasNext": False,
        }
    }


def test_articles_and_comments_come_from_same_fetch():
    articles, comments, restricted = fetch_cafe_articles_and_comments_for_month(
        "https://cafe.naver.com/f-e/cafes/1/menus/1",
        2026,
        8,
        gubun="금호",
        cafe_name="테스트카페",
        group="테스트차종",
        fetch_page=_article_page,
        fetch_comment_page=_comment_page_for_111,
    )
    assert len(articles) == 2
    assert len(comments) == 2  # 댓글 0개짜리 글은 애초에 조회도 안 함
    assert restricted == 0

    # T-04 핵심 승인 기준: 댓글 시트 행수(구분 필터 후) == 게시글 시트 댓글 합계
    gubun_filtered_comment_rows = [c for c in comments if c.gubun == "금호"]
    article_comment_sum = sum(a.comments for a in articles)
    assert len(gubun_filtered_comment_rows) == article_comment_sum == 2


def test_articles_and_comments_excludes_member_only_from_restricted_count():
    # articleId 111은 회원전용(4004)이라 댓글 원문을 못 가져오는 상황을 재현.
    def fetch_comment_page_restricted(cafe_id, article_id, page, **kwargs):
        if article_id == 111:
            raise CommentsRestrictedError(cafe_id, article_id, "카페 멤버만 읽을 수 있는 게시글입니다.")
        return {"result": {"comments": {"items": []}, "hasNext": False}}

    articles, comments, restricted = fetch_cafe_articles_and_comments_for_month(
        "https://cafe.naver.com/f-e/cafes/1/menus/1",
        2026,
        8,
        gubun="금호",
        cafe_name="테스트카페",
        group="테스트차종",
        fetch_page=_article_page,
        fetch_comment_page=fetch_comment_page_restricted,
    )
    assert len(articles) == 2
    assert comments == []
    assert restricted == 2  # articleId 111의 commentCount

    # 호출부가 하는 것과 동일한 검증: 회원전용 제외분을 빼면 정확히 맞아야 한다.
    total_comments_sheet = sum(a.comments for a in articles)
    gubun_filtered = [c for c in comments if c.gubun == "금호"]
    assert len(gubun_filtered) == total_comments_sheet - restricted == 0


def test_fetch_comments_for_article_paginates():
    pages = {
        1: {"result": {"comments": {"items": [
            {"writer": {"nick": "A"}, "content": "1", "updateDate": 1786501400000, "isDeleted": False},
        ]}, "hasNext": True}},
        2: {"result": {"comments": {"items": [
            {"writer": {"nick": "B"}, "content": "2", "updateDate": 1786501500000, "isDeleted": False},
        ]}, "hasNext": False}},
    }

    def fetch(cafe_id, article_id, page, **kwargs):
        return pages[page]

    comments = fetch_comments_for_article(
        "1", 999, gubun="금호", cafe_name="카페", article_title="글", fetch_comment_page=fetch
    )
    assert [c.content for c in comments] == ["1", "2"]


def test_fetch_comments_for_article_skips_deleted():
    def fetch(cafe_id, article_id, page, **kwargs):
        return {
            "result": {
                "comments": {
                    "items": [
                        {"writer": {"nick": "A"}, "content": "살아있음", "updateDate": 1786501400000, "isDeleted": False},
                        {"writer": {"nick": "B"}, "content": "삭제됨", "updateDate": 1786501500000, "isDeleted": True},
                    ]
                },
                "hasNext": False,
            }
        }

    comments = fetch_comments_for_article(
        "1", 999, gubun="금호", cafe_name="카페", article_title="글", fetch_comment_page=fetch
    )
    assert len(comments) == 1
    assert comments[0].content == "살아있음"
