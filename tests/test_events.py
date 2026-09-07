"""T-06 승인 기준 검증: 이벤트 게시글을 카페 로스터 전체에서 정확히 찾아내고,
실제 게시글 댓글 수와 집계가 일치하는지 확인.
"""

from __future__ import annotations

import openpyxl
import pytest

from naver_cafe_crawler.events import (
    EVENT_HEADER,
    collect_event_comments,
    find_event_across_cafes,
    summarize,
    write_event_sheet,
)
from naver_cafe_crawler.naver_api import CommentsRestrictedError

CAFES = [
    {"name": "카니발 포에버", "url": "https://cafe.naver.com/f-e/cafes/1/menus/1", "car_model": "카니발KA4", "board_group": "CRUGEN GT Pro"},
    {"name": "쏘렌토 멤버스", "url": "https://cafe.naver.com/f-e/cafes/2/menus/2", "car_model": "쏘렌토MQ4", "board_group": "CRUGEN GT Pro"},
    {"name": "클럽 싼타페", "url": "https://cafe.naver.com/f-e/cafes/3/menus/3", "car_model": "싼타페TM", "board_group": "CRUGEN GT Pro"},
]


def _article(article_id, subject, ts_ms=1785718777483, read=0, like=0, comment=0):
    return {
        "type": "ARTICLE",
        "item": {
            "articleId": article_id,
            "subject": subject,
            "writeDateTimestamp": ts_ms,
            "readCount": read,
            "likeCount": like,
            "commentCount": comment,
            "writerInfo": {"nickName": "타이어프로", "memberLevelName": "협력업체"},
        },
    }


def fake_fetch_page(cafe_id, menu_id, page, **kwargs):
    if page > 1:
        return {"result": {"articleList": [], "pageInfo": {}}}
    data = {
        "1": [_article(101, "[이벤트] 여름맞이 타이어 이벤트", read=100, like=5, comment=10)],
        "2": [_article(201, "[공지] 여름맞이 타이어 이벤트", read=200, like=8, comment=20)],
        "3": [_article(301, "전혀 다른 글", read=50, like=1, comment=2)],  # 매칭 안 됨
    }
    return {"result": {"articleList": data[cafe_id], "pageInfo": {"lastNavigationPageNumber": 1, "visibleNextButton": False}}}


def test_find_event_across_cafes_matches_normalized_title():
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)
    assert len(results) == 3
    assert results[0].found is True
    assert results[0].views == 100
    assert results[1].found is True  # [공지] 태그만 다르고 정규화하면 같은 제목
    assert results[1].views == 200
    assert results[2].found is False
    assert results[2].views == 0


def test_summarize_totals_only_found_cafes():
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)
    summary = summarize(results, comments=[], restricted=0)
    assert summary.total_views == 300  # 100 + 200, 매칭 안 된 카페(50)는 제외
    assert summary.total_likes == 13
    assert summary.total_comments == 30
    assert summary.cafes_found_in == 2
    assert summary.cafes_total == 3


def test_collect_event_comments_fetches_only_matched_cafes():
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)

    calls = []

    def fake_comment_page(cafe_id, article_id, page, **kwargs):
        calls.append((cafe_id, article_id))
        return {
            "result": {
                "comments": {"items": [
                    {"writer": {"nick": f"참여자{cafe_id}"}, "content": "저요!", "updateDate": 1785718777483, "isDeleted": False},
                ]},
                "hasNext": False,
            }
        }

    comments, restricted = collect_event_comments(results, gubun="금호", fetch_comment_page=fake_comment_page)
    assert restricted == 0
    assert len(comments) == 2  # 매칭된 카페(1, 2)만 조회 -> 각 1개
    assert {cid for cid, aid in calls} == {"1", "2"}
    assert "3" not in {cid for cid, aid in calls}  # 매칭 안 된 카페는 댓글 조회 자체를 안 함


def test_summarize_unique_participants_dedupes_by_writer():
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)
    from naver_cafe_crawler.comments import CommentRecord

    comments = [
        CommentRecord(gubun="금호", cafe="카니발 포에버", article_title="x", write_date="2026.08.10", writer="홍길동", content="1"),
        CommentRecord(gubun="금호", cafe="쏘렌토 멤버스", article_title="x", write_date="2026.08.11", writer="홍길동", content="2"),  # 같은 사람, 다른 카페
        CommentRecord(gubun="금호", cafe="카니발 포에버", article_title="x", write_date="2026.08.10", writer="김철수", content="3"),
    ]
    summary = summarize(results, comments, restricted=0)
    assert summary.unique_participant_count == 2  # 홍길동, 김철수


def test_collect_event_comments_tracks_restricted():
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)

    def restricted_fetch(cafe_id, article_id, page, **kwargs):
        raise CommentsRestrictedError(cafe_id, article_id, "카페 멤버만 읽을 수 있는 게시글입니다.")

    comments, restricted = collect_event_comments(results, gubun="금호", fetch_comment_page=restricted_fetch)
    assert comments == []
    assert restricted == 30  # 10 + 20 (매칭된 두 카페의 commentCount 합)


def test_write_event_sheet_schema(tmp_path):
    results = find_event_across_cafes(CAFES, "여름맞이 타이어 이벤트", 2026, 8, fetch_page=fake_fetch_page)
    summary = summarize(results, comments=[], restricted=0)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "금호"
    write_event_sheet(ws, "여름맞이 타이어 이벤트", results, summary)
    out = tmp_path / "event.xlsx"
    wb.save(out)

    loaded = openpyxl.load_workbook(out, data_only=True)["금호"]
    assert loaded.cell(3, 1).value == 300  # 총조회수
    header_row = [loaded.cell(5, c).value for c in range(1, 1 + len(EVENT_HEADER))]
    assert header_row == EVENT_HEADER
    assert loaded.cell(6, 4).value == "카니발 포에버"
    assert loaded.cell(6, 9).value == "O"
    assert loaded.cell(8, 9).value == "X"
